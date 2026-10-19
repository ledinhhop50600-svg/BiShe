"""生成学校过程材料初稿（开题报告 / 中期检查表）。

用法（项目根目录，需要 python-docx）：
    python tools/make_school_docs.py

产物：
    文件/开题报告（初稿）.docx
    文件/中期检查表（初稿）.docx

说明：
  - 个人信息（姓名/班级学号/指导教师）留占位，由学生填写；
  - 内容与任务书、系统实际实现保持一致，不照搬他人材料；
  - 格式为通用草稿格式，提交前请按学校模板调整字体/表格。
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "文件")
TITLE = "基于SpringBoot+Vue和单智能体的自媒体平台管理系统"


def set_font(run, size=12, bold=False, name="宋体"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def setup(doc):
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.4)


def title(doc, text, size=16):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), size=size, bold=True)
    p.paragraph_format.space_after = Pt(10)


def info_table(doc):
    rows = [
        ("学生姓名", "（待填）", "班级学号", "（待填）"),
        ("指导教师", "（待填）", "专业", "（待填）"),
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            set_font(p.add_run(value), size=10.5, bold=(j % 2 == 0))
    p = doc.add_paragraph()
    set_font(p.add_run("毕业设计（论文）题目："), size=10.5, bold=True)
    set_font(p.add_run(TITLE), size=10.5)
    doc.add_paragraph()


def heading(doc, text):
    p = doc.add_paragraph()
    set_font(p.add_run(text), size=12, bold=True)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def body(doc, text, indent=True):
    p = doc.add_paragraph()
    set_font(p.add_run(text), size=12)
    p.paragraph_format.line_spacing = 1.5
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(2)


def bullet(doc, text):
    p = doc.add_paragraph()
    set_font(p.add_run(text), size=12)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(2)


def make_open_report():
    doc = Document()
    setup(doc)
    title(doc, "广州航海学院毕业设计（论文）开题报告")
    title(doc, TITLE, size=14)
    info_table(doc)

    heading(doc, "一、选题的目的和意义")
    body(doc,
         "随着自媒体内容在公众号、小红书、微博、B站、抖音等平台之间分发，创作者需要把同一篇内容改写成"
         "符合各平台格式与社区规范的版本；平台运营方则需要判断内容是否合规。人工完成这两件事成本高、"
         "标准不一致，而且不同审核员的尺度会有差异。通用大模型能够理解语境，但通常不了解平台专有规则"
         "与法规条款，给出的结论也缺少可追溯的依据。")
    body(doc,
         "本课题拟设计并实现一个面向多平台的内容创作、审核与适配系统：以 SpringBoot + Vue 实现管理主体，"
         "以独立的 Python 单智能体服务提供发布前预检、正式初审与跨平台内容适配，审核员保留最终决定权。"
         "研究目标不是让模型替代人，而是把模型放进一个有约束、可追溯、可纠正的业务流程中。")
    body(doc,
         "意义主要有三点：一是应用价值，降低审核与跨平台改写的人工成本，给创作者具体可执行的修改建议；"
         "二是技术价值，验证“规则库 + 检索增强生成（RAG） + 结构化输出 + 人机协同”在内容合规场景中的"
         "可行性，并用可复现的对比实验说明各方案的差异；三是工程价值，完整走通前后端分离、独立智能体服务、"
         "接口契约、降级与留痕等工程环节，为本人从事 Agent 应用开发积累实践。")

    heading(doc, "二、国内外研究现状及存在的问题")
    body(doc,
         "内容审核方面，早期方案以关键词与敏感词匹配为主，实现简单但无法理解语境，容易把作品名、普法引用、"
         "新闻转述误判为违规；随后出现基于传统机器学习的分类器与平台自建审核系统；近年来大模型被广泛用于"
         "内容安全，公开数据集如 ToxiCN（ACL 2023）与 COLD（EMNLP 2022）为中文有害内容检测提供了评测基础。"
         "Markov 等（AAAI 2023）也讨论了真实业务中不良内容检测的整体方法。")
    body(doc,
         "检索增强生成（RAG）由 Lewis 等（NeurIPS 2020）提出，通过引入外部知识减少模型幻觉，已被用于"
         "知识密集型的问答与合规场景。Agent 方面，ReAct（ICLR 2023）等工作把推理与工具调用结合起来；"
         "在边界清晰的任务上，单 Agent 往往比多 Agent 更容易控制成本与调试。人机协同（human-in-the-loop）"
         "则强调由人做最终决策，并把人的改判沉淀为数据。")
    body(doc, "结合本课题，现有做法主要存在以下问题：")
    bullet(doc, "（1）通用模型不了解平台专有规则与法规条款，判定理由无法追溯到具体规则；")
    bullet(doc, "（2）关键词基线语境误判多，纯大模型结论不稳定、存在幻觉，成本也缺少量化；")
    bullet(doc, "（3）各平台标题/正文字数、标签、形态差异大，人工逐个改写费时，格式约束缺少程序化校验；")
    bullet(doc, "（4）AI 结论缺少人工复核与改判留痕，难以评估“AI 到底帮了多少忙”；")
    bullet(doc, "（5）个人开发者无法获得各平台发布接口，真实“一键发布”不可行，需要明确系统边界。")

    heading(doc, "三、主要研究内容")
    body(doc, "（1）需求分析与总体架构：三类角色（管理员/审核员/创作者）、前后端分离、智能体独立服务、"
              "审核与适配接口契约。")
    body(doc, "（2）基础管理功能：用户、角色、菜单的 RBAC 权限体系，稿件分类与标签、举报处理、敏感词库"
              "（作为关键词基线）。")
    body(doc, "（3）内容与审核流程：稿件状态机（草稿/待审/发布/驳回/阻断/下架）、投稿与审核工作台、"
              "人机协同（AI 初判 + 审核员采纳/改判）、改判原因留痕、内容哈希判断结论是否过期。")
    body(doc, "（4）审核智能体：六类规则库、四级风险分级（通过/疑似/违规/严重）、RAG 混合检索、长文分段"
              "与关键词定位、结构化 JSON 输出（等级/理由/命中规则/违规片段/建议）、模型不可用时的降级。")
    body(doc, "（5）跨平台内容适配：按平台约束与形态改写（小红书短笔记、公众号长文、微博短内容、"
              "B站视频简介、抖音口播脚本）；字数与标签由程序做格式自检，适配结果复用审核规则做平台合规预检，"
              "并按“稿件 + 平台”留存版本。")
    body(doc, "（6）评测与实验：构建 52 条审核评测集（自建分层样本 + ToxiCN 公开样本），对比关键词规则、"
              "纯大模型、RAG+大模型、全量规则四种方案的准确率/F1/成本/耗时；对适配结果统计格式满足率、"
              "长文压缩比，并进行 LLM 评分与人工评分。")
    body(doc, "（7）系统测试与交付：自动化验证脚本、一键启动、运行说明、答辩材料。")

    heading(doc, "四、研究方法、步骤和措施")
    body(doc, "研究方法：文献调研与需求分析、原型迭代开发、对比实验、人工评估。")
    body(doc, "研究步骤按依赖关系组织（不按日历时间）：")
    bullet(doc, "① 需求澄清与关键决策记录；② 技术选型与最小可运行验证；③ 数据库与 RBAC；")
    bullet(doc, "④ 内容管理、投稿与审核流程；⑤ 审核智能体（规则/RAG/提示词/降级）；")
    bullet(doc, "⑥ 人机协同闭环（初审、采纳/改判、反馈、过期）；⑦ 跨平台适配；")
    bullet(doc, "⑧ 评测集与对比实验；⑨ 系统测试、论文与答辩材料。")
    body(doc, "主要措施：")
    bullet(doc, "（1）接口契约先行：审核/适配的请求与响应字段、错误码冻结后再并行开发；")
    bullet(doc, "（2）规则可溯源：每条规则保留法规/平台规范来源，判定理由必须引用规则编号；")
    bullet(doc, "（3）客观约束交给程序：字数、标签数量与格式由程序校验，模型只负责内容创作与语义判断；")
    bullet(doc, "（4）失败可降级：智能体不可用时不阻断业务，写失败记录并转人工；")
    bullet(doc, "（5）成本可控：使用轻量模型、关闭思考模式、回放缓存，记录每次 token 与耗时；")
    bullet(doc, "（6）全链路留痕：每次判定、采纳/改判、适配版本都可追溯，作为评测数据来源。")
    body(doc, "预期成果：可运行的前后端 + 智能体系统、毕业设计论文、可复现的实验数据与演示材料。")

    heading(doc, "五、主要参考资料")
    refs = [
        "[1] 国家互联网信息办公室. 网络信息内容生态治理规定[Z]. 2019.",
        "[2] 全国人民代表大会常务委员会. 中华人民共和国广告法（2021年修正）[Z]. 2021.",
        "[3] Lewis P, Perez E, Piktus A, et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks[C]. NeurIPS, 2020.",
        "[4] Yao S, Zhao J, Yu D, et al. ReAct: Synergizing Reasoning and Acting in Language Models[C]. ICLR, 2023.",
        "[5] Wei J, Wang X, Schuurmans D, et al. Chain-of-Thought Prompting Elicits Reasoning in Large Language Models[C]. NeurIPS, 2022.",
        "[6] Markov T, Zhang C, Agarwal S, et al. A Holistic Approach to Undesired Content Detection in the Real World[C]. AAAI, 2023.",
        "[7] ToxiCN: A Benchmark for Chinese Toxic Content Detection[C]. ACL, 2023.",
        "[8] COLD: A Benchmark for Chinese Offensive Language Detection[C]. EMNLP, 2022.",
        "[9] DeepSeek. DeepSeek API 文档[EB/OL]. https://api-docs.deepseek.com.",
        "[10] Vue.js 3 官方文档[EB/OL]. https://cn.vuejs.org.",
        "[11] Spring Boot Reference Documentation[EB/OL]. https://docs.spring.io/spring-boot.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        set_font(p.add_run(ref), size=10.5)
        p.paragraph_format.line_spacing = 1.3

    heading(doc, "六、指导教师意见")
    for _ in range(4):
        doc.add_paragraph()
    body(doc, "指导教师签字：______________", indent=False)
    body(doc, "年    月    日", indent=False)

    path = os.path.join(OUT_DIR, "开题报告（初稿）.docx")
    doc.save(path)
    return path


def make_midterm():
    doc = Document()
    setup(doc)
    title(doc, "广州航海学院毕业设计（论文）中期检查表")
    title(doc, TITLE, size=13)
    info_table(doc)

    heading(doc, "一、毕业设计的主要工作内容和计划进度（由学生填写）")
    body(doc,
         "本课题拟完成一个面向多平台的内容创作、审核与适配系统：SpringBoot + Vue 实现管理主体，"
         "Python 单智能体服务提供发布前预检、正式初审与跨平台适配；审核员保留终审与改判权；"
         "同时构建审核评测集并完成对比实验。")
    body(doc,
         "计划进度：需求与方案设计 → 数据库与权限 → 内容管理与投稿审核 → 审核智能体 → 人机协同 → "
         "跨平台适配 → 评测与实验 → 系统测试与论文。")

    heading(doc, "二、目前已完成工作情况（由学生填写）")
    bullet(doc, "（1）完成需求分析、数据库设计（16 张表）、技术选型与关键决策记录；")
    bullet(doc, "（2）完成 RBAC 三角色权限、动态菜单、用户/角色/菜单管理；")
    bullet(doc, "（3）完成稿件状态机、投稿、审核工作台、举报处理、敏感词库；")
    bullet(doc, "（4）完成审核智能体：规则库、四级分级、RAG 混合检索、长文定位、结构化输出与降级；")
    bullet(doc, "（5）完成人机协同闭环：AI 初审、采纳/改判、改判原因留痕、创作者反馈、内容哈希过期；")
    bullet(doc, "（6）完成跨平台适配：多平台改写、程序格式自检、平台合规预检、版本留存；")
    bullet(doc, "（7）完成 52 条审核评测集、四方案对比实验、适配质量评估与系统内评测中心；")
    bullet(doc, "（8）完成注册/个人中心、已发布稿件改后重审、自动化验证脚本与一键启动。")
    body(doc, "目前主要功能与实验已完成，正在补充测试记录、论文撰写与答辩材料。")

    heading(doc, "三、存在的主要问题和解决方案（由学生填写）")
    bullet(doc, "（1）平台规则与字数限制会改版：约束放在数据库中，管理员可维护；数值核对日期与来源写入文档。")
    bullet(doc, "（2）长文适配偶发超限：提示词按上限的 85% 控制，程序格式自检对超限给出明确警告，不静默通过。")
    bullet(doc, "（3）模型成本与稳定性：选用轻量模型、关闭思考模式、回放缓存；不可用时降级为人工，不阻断业务。")
    bullet(doc, "（4）评测样本量有限、人工标注一致性一般：引入公开数据集，计算 Cohen's κ，论文中如实说明局限。")
    bullet(doc, "（5）个人开发者无法接入各平台发布接口：明确系统边界，把价值放在适配与合规预检，不承诺真实发布。")

    heading(doc, "四、指导教师对学生工作态度评价")
    body(doc, "□认真    □较认真    □一般    □差", indent=False)
    heading(doc, "五、指导教师对学生完成工作质量评价")
    body(doc, "□优    □良    □中    □一般    □差", indent=False)
    heading(doc, "六、指导教师意见和建议")
    for _ in range(3):
        doc.add_paragraph()
    body(doc, "指导教师签字：______________", indent=False)
    body(doc, "年    月    日", indent=False)

    path = os.path.join(OUT_DIR, "中期检查表（初稿）.docx")
    doc.save(path)
    return path


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    p1 = make_open_report()
    p2 = make_midterm()
    print("已生成：", os.path.relpath(p1, ROOT))
    print("已生成：", os.path.relpath(p2, ROOT))
    print("个人信息（姓名/班级学号/指导教师/专业）请按学校要求填写；格式提交前按模板微调。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
